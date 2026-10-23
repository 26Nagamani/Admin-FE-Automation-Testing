from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from datetime import datetime
import os
import json
from pathlib import Path


class ExcelService:

    ROOT_FOLDER = r"D:\AutomationReports"

    RESULT_HEADERS = [
        "Test ID",
        "Module Name",
        "Data Test ID",
        "Result",
        "Matched Fields",
        "Matched Documents",
        "Error Message",
        "Backend URL Hit",
        "API Response",
        "Duration",
        "Payload",
        "Screenshot Path",
        # "Video Path",
        "Execution Time",
    ]

    @staticmethod
    def _current_pytest_test_name():
        current_test = os.environ.get("PYTEST_CURRENT_TEST", "")
        if not current_test:
            return ""

        test_id = current_test.split(" ", 1)[0]
        test_name = test_id.rsplit("::", 1)[-1]
        return test_name.replace("/", "_").replace("\\", "_")

    @staticmethod
    def _current_screenshot_path():
        test_name = ExcelService._current_pytest_test_name()
        if not test_name:
            return ""

        screenshot_path = Path("reports/screenshots") / f"{test_name}.png"
        return str(screenshot_path)

    @staticmethod
    def _find_result_row(sheet, headers, test_id, data_test_id):
        test_id_column = headers.index("Test ID") + 1
        data_test_id_column = headers.index("Data Test ID") + 1

        for row_number in range(2, sheet.max_row + 1):
            row_test_id = str(sheet.cell(row=row_number, column=test_id_column).value or "")
            row_data_test_id = str(sheet.cell(row=row_number, column=data_test_id_column).value or "")
            if row_test_id == str(test_id) and row_data_test_id == str(data_test_id):
                return row_number

        return sheet.max_row + 1

    @staticmethod
    def _ensure_headers(sheet):
        existing_headers = [cell.value for cell in sheet[1]]

        if not any(existing_headers):
            for column_index, header in enumerate(ExcelService.RESULT_HEADERS, start=1):
                sheet.cell(row=1, column=column_index).value = header
            return ExcelService.RESULT_HEADERS

        for header in ExcelService.RESULT_HEADERS:
            if header not in existing_headers:
                sheet.cell(row=1, column=len(existing_headers) + 1).value = header
                existing_headers.append(header)

        return existing_headers

    @staticmethod
    def _format_report(sheet, headers, target_row=None):
        header_fill = PatternFill("solid", fgColor="1F4E78")
        pass_fill = PatternFill("solid", fgColor="C6EFCE")
        fail_fill = PatternFill("solid", fgColor="FFC7CE")
        pass_font = Font(color="006100", bold=True)
        fail_font = Font(color="9C0006", bold=True)

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        widths = {
            "Test ID": 16,
            "Module Name": 20,
            "Data Test ID": 28,
            "Result": 12,
            "Matched Fields": 30,
            "Matched Documents": 34,
            "Error Message": 45,
            "Backend URL Hit": 55,
            "API Response": 70,
            "Duration": 12,
            "Payload": 40,
            "Screenshot Path": 38,
            "Execution Time": 22,
        }
        for column_index, header in enumerate(headers, start=1):
            sheet.column_dimensions[sheet.cell(1, column_index).column_letter].width = widths.get(
                header, 20
            )

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 30

        if target_row is None:
            return

        sheet.row_dimensions[target_row].height = 72
        for cell in sheet[target_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        result_column = headers.index("Result") + 1
        result_cell = sheet.cell(target_row, result_column)
        result_cell.alignment = Alignment(horizontal="center", vertical="center")
        if str(result_cell.value).upper() == "PASS":
            for cell in sheet[target_row]:
                cell.fill = pass_fill
            result_cell.font = pass_font
        elif str(result_cell.value).upper() == "FAIL":
            for cell in sheet[target_row]:
                cell.fill = fail_fill
            result_cell.font = fail_font

    @staticmethod
    def clear_module_results(module_name):
        current_date = datetime.now().strftime("%Y-%m-%d")
        module_folder = os.path.join(
            ExcelService.ROOT_FOLDER,
            module_name
        )
        excel_file = os.path.join(
            module_folder,
            f"{current_date}.xlsx"
        )

        os.makedirs(module_folder, exist_ok=True)

        workbook = Workbook()
        sheet = workbook.active
        for column_index, header in enumerate(ExcelService.RESULT_HEADERS, start=1):
            sheet.cell(row=1, column=column_index).value = header

        ExcelService._format_report(sheet, ExcelService.RESULT_HEADERS)

        print(f"[ExcelService] Clearing previous results: {excel_file}")
        workbook.save(excel_file)
        print("[ExcelService] Previous results cleared")

    @staticmethod
    def write_result(
        module_name,
        test_id,
        data_test_id,
        result,
        error_message="",
        backend_url="",
        api_response="",
        matched_fields="",
        matched_documents="",
        duration="",
        payload="",
        screenshot_path="",
        # video_path=""
    ):

        # Convert payload dict → JSON string (important for readability)
        if isinstance(payload, dict):
            payload = json.dumps(payload, indent=2)

        current_date = datetime.now().strftime("%Y-%m-%d")

        module_folder = os.path.join(
            ExcelService.ROOT_FOLDER,
            module_name
        )

        os.makedirs(module_folder, exist_ok=True)

        excel_file = os.path.join(
            module_folder,
            f"{current_date}.xlsx"
        )

        # Load or create workbook
        if os.path.exists(excel_file):
            workbook = load_workbook(excel_file)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active

        headers = ExcelService._ensure_headers(sheet)

        if str(result).upper() == "FAIL":
            screenshot_path = screenshot_path or ExcelService._current_screenshot_path()
        else:
            screenshot_path = ""
        # video_path = ""

        row_values = {
            "Test ID": test_id,
            "Module Name": module_name,
            "Data Test ID": data_test_id,
            "Result": result,
            "Matched Fields": matched_fields,
            "Matched Documents": matched_documents,
            "Error Message": error_message,
            "Backend URL Hit": backend_url,
            "API Response": api_response,
            "Duration": duration,
            "Payload": payload,
            "Screenshot Path": screenshot_path,
            # "Video Path": video_path,
            "Execution Time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }

        # Upsert result row for this test case instead of appending duplicates.
        target_row = ExcelService._find_result_row(sheet, headers, test_id, data_test_id)
        for column_index, header in enumerate(headers, start=1):
            sheet.cell(row=target_row, column=column_index).value = row_values.get(header, "")

        ExcelService._format_report(sheet, headers, target_row)

        print(f"[ExcelService] Saving to: {excel_file}")
        workbook.save(excel_file)
        print(f"[ExcelService] Saved successfully") 

    @staticmethod
    def update_data_row(file_path, sheet_name, match_column, match_value, updates):
        workbook = load_workbook(file_path)
        sheet = workbook[sheet_name]

        headers = [cell.value for cell in sheet[1]]

        for column_name in updates:
            if column_name not in headers:
                sheet.cell(row=1, column=len(headers) + 1).value = column_name
                headers.append(column_name)

        if match_column not in headers:
            raise ValueError(f"Column not found: {match_column}")

        match_column_index = headers.index(match_column) + 1

        target_row = None
        for row_number in range(2, sheet.max_row + 1):
            if str(sheet.cell(row=row_number, column=match_column_index).value or "") == str(match_value):
                target_row = row_number
                break

        if target_row is None:
            raise ValueError(f"No row found for {match_column}={match_value}")

        for column_name, value in updates.items():
            column_index = headers.index(column_name) + 1
            sheet.cell(row=target_row, column=column_index).value = value

        workbook.save(file_path)
