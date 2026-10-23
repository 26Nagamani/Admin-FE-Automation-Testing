from openpyxl import load_workbook
import warnings

def read_excel(file_path, sheet_name):
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style, apply openpyxl's default",
            category=UserWarning,
            module="openpyxl.styles.stylesheet",
        )
        workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    data = []

    headers = [cell.value for cell in sheet[1]]

    for row in sheet.iter_rows(min_row=2, values_only=True):
        data.append(dict(zip(headers, row)))

    return data
